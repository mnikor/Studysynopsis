from __future__ import annotations

import base64
import json
import os
import re
import zipfile
from io import BytesIO
from typing import Any

import httpx
from fastapi import FastAPI, HTTPException, Response
from fastapi.middleware.cors import CORSMiddleware

app = FastAPI(title="Study Synopsis Studio API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:3001",
        "http://127.0.0.1:3001",
        "http://localhost:5173",
        "http://127.0.0.1:5173",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

OPENAI_MODEL = os.getenv("OPENAI_MODEL", "gpt-5")
OPENAI_API_URL = os.getenv("OPENAI_API_URL", "https://api.openai.com/v1/responses")


def string_schema() -> dict[str, Any]:
    return {"type": "string"}


def array_schema(items: dict[str, Any] | None = None) -> dict[str, Any]:
    return {"type": "array", "items": items or string_schema()}


def object_schema(properties: dict[str, Any]) -> dict[str, Any]:
    return {
        "type": "object",
        "properties": properties,
        "required": list(properties.keys()),
        "additionalProperties": False,
    }


STUDY_FIELDS = [
    "studyTitle",
    "sponsor",
    "category",
    "subcategory",
    "developmentStage",
    "protocolGuardrailProfile",
    "phase3SafetyDataAvailable",
    "protocolGuardrailNotes",
    "primaryEvidenceUseIntent",
    "secondaryEvidenceUseIntents",
    "primaryStrategicObjective",
    "secondaryStrategicObjectives",
    "customStrategicObjective",
    "therapeuticArea",
    "disease",
    "customDisease",
    "topIntervention",
    "topInterventionClass",
    "topComparator",
    "topLineOfTherapy",
    "selectedEndpoints",
    "customEndpoints",
    "indication",
    "primaryObjective",
    "secondaryObjectives",
    "designOverview",
    "population",
    "intervention",
    "comparator",
    "outcomes",
    "timeline",
    "geography",
    "sampleSize",
    "eligibility",
    "operationalNotes",
]


def study_output_schema(extra: dict[str, Any] | None = None) -> dict[str, Any]:
    arrays = {"secondaryEvidenceUseIntents", "secondaryStrategicObjectives", "selectedEndpoints"}
    properties = {field: array_schema() if field in arrays else string_schema() for field in STUDY_FIELDS}
    properties.update(extra or {})
    return object_schema(properties)


PICO_SCHEMA = object_schema(
    {
        "population": string_schema(),
        "intervention": string_schema(),
        "comparator": string_schema(),
        "outcomes": string_schema(),
        "timeframe": string_schema(),
        "prompt": string_schema(),
    }
)

STATS_SCHEMA = object_schema(
    {
        "estimand": string_schema(),
        "endpointType": string_schema(),
        "hypothesis": string_schema(),
        "alpha": string_schema(),
        "power": string_schema(),
        "effectSize": string_schema(),
        "variability": string_schema(),
        "attrition": string_schema(),
        "analysisModel": string_schema(),
        "missingData": string_schema(),
        "rationale": string_schema(),
        "prompt": string_schema(),
    }
)

LITERATURE_SCHEMA = object_schema(
    {
        "researchQuestion": string_schema(),
        "databases": string_schema(),
        "evidenceWindow": string_schema(),
        "inclusionCriteria": string_schema(),
        "exclusionCriteria": string_schema(),
        "keywords": string_schema(),
        "pubmedQuery": string_schema(),
        "embaseQuery": string_schema(),
        "greyLiteraturePlan": string_schema(),
        "backgroundThemes": string_schema(),
        "prompt": string_schema(),
    }
)

ALIGNMENT_SCHEMA = object_schema(
    {
        "status": {"type": "string", "enum": ["aligned", "partially_aligned", "conflicts"]},
        "summary": string_schema(),
        "concerns": array_schema(),
        "recommendations": array_schema(),
    }
)

OBJECTIVE_OPTION_SCHEMA = object_schema(
    {
        "label": string_schema(),
        "positioning": string_schema(),
        "primaryObjective": string_schema(),
        "secondaryObjectives": array_schema(),
        "designOverview": string_schema(),
        "alignment": ALIGNMENT_SCHEMA,
    }
)

ENDPOINT_OPTION_SCHEMA = object_schema(
    {
        "label": string_schema(),
        "positioning": string_schema(),
        "primaryEndpoint": string_schema(),
        "secondaryEndpoints": array_schema(),
        "exploratoryEndpoints": array_schema(),
        "rationale": string_schema(),
        "alignment": ALIGNMENT_SCHEMA,
    }
)

SCHEDULE_COLUMN_SCHEMA = object_schema(
    {
        "id": string_schema(),
        "phase": string_schema(),
        "period": string_schema(),
        "visit": string_schema(),
        "footnote": string_schema(),
    }
)

SCHEDULE_CELL_SCHEMA = object_schema({"columnId": string_schema(), "value": string_schema()})
SCHEDULE_ROW_SCHEMA = object_schema(
    {
        "id": string_schema(),
        "group": string_schema(),
        "activity": string_schema(),
        "cells": array_schema(SCHEDULE_CELL_SCHEMA),
        "notes": string_schema(),
    }
)

IMPACT_DOMAIN_SCHEMA = object_schema(
    {
        "id": string_schema(),
        "label": string_schema(),
        "score": {"type": "number"},
        "impact": {"type": "string", "enum": ["low", "medium", "high"]},
        "rationale": string_schema(),
    }
)

COMPLEXITY_DRIVER_SCHEMA = object_schema(
    {
        "label": string_schema(),
        "impact": {"type": "string", "enum": ["low", "medium", "high"]},
        "rationale": string_schema(),
    }
)

TRADEOFF_RECOMMENDATION_SCHEMA = object_schema(
    {
        "id": string_schema(),
        "target": string_schema(),
        "category": string_schema(),
        "recommendation": string_schema(),
        "expectedBenefit": string_schema(),
        "dataQualityRisk": {"type": "string", "enum": ["low", "medium", "high"]},
        "endpointProtection": string_schema(),
        "rationale": string_schema(),
    }
)

STUDY_SCHEMA_LANE = object_schema({"id": string_schema(), "label": string_schema(), "description": string_schema()})
STUDY_SCHEMA_NODE = object_schema(
    {
        "id": string_schema(),
        "laneId": string_schema(),
        "label": string_schema(),
        "subtitle": string_schema(),
        "kind": {
            "type": "string",
            "enum": ["start", "screening", "decision", "arm", "cohort", "database", "assessment", "analysis", "output", "milestone"],
        },
        "column": {"type": "integer"},
        "row": {"type": "integer"},
    }
)
STUDY_SCHEMA_EDGE = object_schema(
    {
        "id": string_schema(),
        "from": string_schema(),
        "to": string_schema(),
        "label": string_schema(),
        "style": {"type": "string", "enum": ["solid", "dashed"]},
    }
)


ACTION_CONFIG: dict[str, dict[str, Any]] = {
    "suggest_objectives": {
        "schema_name": "suggested_study_objectives",
        "schema": object_schema({"guidance": string_schema(), "options": array_schema(OBJECTIVE_OPTION_SCHEMA)}),
        "prompt": (
            "You are a senior clinical strategist. Propose exactly 3 differentiated research objective packages. "
            "Objectives must be scientific and testable, not strategic phrases such as support HTA or label change. "
            "Use evidence intent to shape comparator rigor and endpoint discipline. Keep secondary objectives focused. "
            "If the user's request conflicts with the study context, mark the alignment as partially_aligned or conflicts."
        ),
        "tokens": 9000,
    },
    "suggest_endpoints": {
        "schema_name": "suggested_study_endpoints",
        "schema": object_schema({"guidance": string_schema(), "options": array_schema(ENDPOINT_OPTION_SCHEMA)}),
        "prompt": (
            "You are a clinical endpoint specialist. Propose exactly 3 endpoint packages. "
            "Avoid endpoint laundry lists. In most cases include 1 primary endpoint, 1 to 2 key secondary endpoints, "
            "and at most 1 exploratory endpoint only if clearly justified. Include PRO, PK/PD, biomarkers, biospecimens, "
            "special imaging, or extra assessments only when decision-critical."
        ),
        "tokens": 8000,
    },
    "suggest_population": {
        "schema_name": "suggested_population",
        "schema": object_schema({"population": string_schema(), "eligibility": string_schema(), "rationale": string_schema()}),
        "prompt": (
            "Draft a concise target population and eligibility highlights from the study objective, disease, intervention, "
            "comparator, endpoints, study type, and evidence intent. Avoid unsupported highly specific criteria."
        ),
        "tokens": 5000,
    },
    "assess_impact": {
        "schema_name": "impact_assessment",
        "schema": object_schema(
            {
                "executiveSummary": string_schema(),
                "domains": array_schema(IMPACT_DOMAIN_SCHEMA),
                "blockers": array_schema(),
                "strengthenActions": array_schema(),
            }
        ),
        "prompt": (
            "Assess likely evidence impact as an early design pressure-test. Return exactly 5 domains in order: "
            "guideline, label, publication, practice, hta. Be conservative and decision-oriented."
        ),
        "tokens": 6000,
    },
    "generate_study_schema": {
        "schema_name": "study_schema",
        "schema": object_schema(
            {
                "title": string_schema(),
                "schemaType": string_schema(),
                "orientation": {"type": "string", "enum": ["horizontal", "vertical"]},
                "detailLevel": {"type": "string", "enum": ["simple", "standard", "detailed"]},
                "lanes": array_schema(STUDY_SCHEMA_LANE),
                "nodes": array_schema(STUDY_SCHEMA_NODE),
                "edges": array_schema(STUDY_SCHEMA_EDGE),
                "notes": array_schema(),
            }
        ),
        "prompt": (
            "Create structured data for a one-slide executive study schema. Keep it readable and concise. "
            "Return nodes, lanes, and edges only; do not return SVG or prose."
        ),
        "tokens": 7000,
    },
    "extract_pico_stats": {
        "schema_name": "pico_stats",
        "schema": object_schema({"pico": PICO_SCHEMA, "stats": STATS_SCHEMA}),
        "prompt": (
            "Extract clean PICO elements and draft first-pass statistical assumptions from the study design. "
            "Keep outputs concise, specific, and suitable for synopsis drafting."
        ),
        "tokens": 9000,
    },
    "build_literature": {
        "schema_name": "literature_plan",
        "schema": object_schema({"literature": LITERATURE_SCHEMA}),
        "prompt": (
            "Build a pragmatic literature plan that supports the synopsis rationale. Keep it focused on the decision "
            "the study should support, not a broad unfocused review."
        ),
        "tokens": 7000,
    },
    "generate_schedule": {
        "schema_name": "study_schedule_of_activities",
        "schema": object_schema({"purpose": string_schema(), "prompt": string_schema(), "columns": array_schema(SCHEDULE_COLUMN_SCHEMA), "rows": array_schema(SCHEDULE_ROW_SCHEMA)}),
        "prompt": (
            "Generate a Schedule of Activities as structured data. Columns must preserve phase, period, and visit hierarchy. "
            "Rows must be grouped. Use X for required activities and blank for not required. Support multiple arms, cohorts, "
            "or dosing schedules using concise conditional rows or separate treatment groups."
        ),
        "tokens": 12000,
    },
    "analyze_schedule": {
        "schema_name": "schedule_insights",
        "schema": object_schema(
            {
                "complexity": object_schema(
                    {
                        "overallLevel": {"type": "string", "enum": ["Low", "Moderate", "High", "Very High"]},
                        "overallScore": {"type": "number"},
                        "executiveSummary": string_schema(),
                        "visitBurdenScore": {"type": "number"},
                        "assessmentBurdenScore": {"type": "number"},
                        "operationalBurdenScore": {"type": "number"},
                        "drivers": array_schema(COMPLEXITY_DRIVER_SCHEMA),
                        "mitigations": array_schema(),
                    }
                ),
                "tradeoff": object_schema(
                    {
                        "executiveSummary": string_schema(),
                        "protectedElements": array_schema(),
                        "recommendations": array_schema(TRADEOFF_RECOMMENDATION_SCHEMA),
                        "safeguards": array_schema(),
                    }
                ),
            }
        ),
        "prompt": (
            "Review SoA complexity and data-value trade-offs. Identify lower-risk simplifications without weakening objectives, "
            "primary endpoint support, essential safety, or feasibility."
        ),
        "tokens": 9000,
    },
    "draft_sections": {
        "schema_name": "draft_sections",
        "schema": object_schema({"sections": array_schema(object_schema({"id": string_schema(), "title": string_schema(), "body": string_schema()}))}),
        "prompt": "Draft concise synopsis section bodies for each included section using all provided study context.",
        "tokens": 9000,
    },
    "generate_synopsis": {
        "schema_name": "final_synopsis",
        "schema": object_schema({"sections": array_schema(object_schema({"id": string_schema(), "title": string_schema(), "body": string_schema()}))}),
        "prompt": "Generate polished final study synopsis section bodies using the provided section configuration and study context.",
        "tokens": 12000,
    },
    "import_study_document": {
        "schema_name": "study_document_import",
        "schema": object_schema(
            {
                "summary": string_schema(),
                "unresolved": array_schema(),
                "assumptions": array_schema(),
                "study": study_output_schema({"endpointCandidates": array_schema()}),
            }
        ),
        "prompt": (
            "Extract structured study-design information from the uploaded source. Use blank strings and empty arrays for "
            "unavailable fields. Do not force weak matches to controlled lists."
        ),
        "tokens": 9000,
    },
}


def decode_data_url(data_url: str) -> tuple[str, bytes]:
    match = re.match(r"^data:([^;]+);base64,(.+)$", data_url)
    if not match:
        raise HTTPException(status_code=400, detail="The uploaded file could not be decoded.")
    return match.group(1), base64.b64decode(match.group(2))


def decode_xml_entities(value: str) -> str:
    return (
        value.replace("&amp;", "&")
        .replace("&lt;", "<")
        .replace("&gt;", ">")
        .replace("&quot;", '"')
        .replace("&#39;", "'")
        .replace("&#xA;", "\n")
    )


def extract_office_text(buffer: bytes, kind: str) -> str:
    text_chunks: list[str] = []
    with zipfile.ZipFile(BytesIO(buffer)) as archive:
      names = archive.namelist()
      if kind == "docx":
          targets = ["word/document.xml", *[name for name in names if re.match(r"^word/(header|footer)\d+\.xml$", name)]]
          text_pattern = re.compile(r"<w:t[^>]*>([\s\S]*?)</w:t>")
          paragraph_pattern = re.compile(r"</w:p>")
      else:
          targets = sorted(
              [name for name in names if re.match(r"^ppt/slides/slide\d+\.xml$", name)],
              key=lambda name: int(re.search(r"slide(\d+)\.xml$", name).group(1)),
          )
          text_pattern = re.compile(r"<a:t>([\s\S]*?)</a:t>")
          paragraph_pattern = re.compile(r"</a:p>")

      for name in targets:
          if name not in names:
              continue
          xml = archive.read(name).decode("utf-8", errors="ignore")
          xml = paragraph_pattern.sub("\n", xml)
          text = " ".join(decode_xml_entities(match.group(1)) for match in text_pattern.finditer(xml))
          text = re.sub(r"\s+\n", "\n", text)
          text = re.sub(r"\n\s+", "\n", text)
          text = re.sub(r"[ \t]{2,}", " ", text).strip()
          if text:
              text_chunks.append(text)
    return "\n\n".join(text_chunks).strip()


def extract_response_text(response_json: dict[str, Any]) -> str:
    if isinstance(response_json.get("output_text"), str):
        return response_json["output_text"]
    chunks: list[str] = []
    for item in response_json.get("output", []):
        for content in item.get("content", []) if isinstance(item, dict) else []:
            if isinstance(content, dict):
                text = content.get("text") or content.get("output_text")
                if isinstance(text, str):
                    chunks.append(text)
    return "\n".join(chunks).strip()


async def call_openai_structured(action: str, payload: dict[str, Any], user_content: Any | None = None) -> Any:
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        raise HTTPException(status_code=500, detail="OPENAI_API_KEY is not configured on the server.")

    config = ACTION_CONFIG[action]
    request_body = {
        "model": OPENAI_MODEL,
        "reasoning": {"effort": "low"},
        "max_output_tokens": config["tokens"],
        "input": [
            {"role": "system", "content": config["prompt"]},
            {"role": "user", "content": user_content if user_content is not None else json.dumps(payload, indent=2)},
        ],
        "text": {
            "format": {
                "type": "json_schema",
                "name": config["schema_name"],
                "schema": config["schema"],
                "strict": True,
            }
        },
    }

    async with httpx.AsyncClient(timeout=180) as client:
        response = await client.post(
            OPENAI_API_URL,
            headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
            json=request_body,
        )

    if response.status_code >= 400:
        raise HTTPException(status_code=response.status_code, detail=response.text)

    response_json = response.json()
    text = extract_response_text(response_json)
    if not text:
        raise HTTPException(status_code=502, detail=f"The model did not return structured text. Status: {response_json.get('status')}")
    try:
        return json.loads(text)
    except json.JSONDecodeError as exc:
        raise HTTPException(status_code=502, detail=f"The model returned invalid JSON: {text[:300]}") from exc


@app.get("/health")
async def health() -> dict[str, str]:
    return {"status": "ok"}


@app.post("/api/openai")
async def openai_endpoint(payload: dict[str, Any]) -> Any:
    action = payload.get("action")
    if action not in ACTION_CONFIG:
        raise HTTPException(status_code=400, detail="Unsupported action.")

    if action == "import_study_document":
        file_name = str(payload.get("fileName", ""))
        mime_type = str(payload.get("mimeType", ""))
        file_data = str(payload.get("fileData", ""))
        lower_name = file_name.lower()
        is_pdf = mime_type == "application/pdf" or lower_name.endswith(".pdf")
        is_docx = lower_name.endswith(".docx") or "wordprocessingml" in mime_type
        is_pptx = lower_name.endswith(".pptx") or "presentationml" in mime_type

        if not (is_pdf or is_docx or is_pptx):
            raise HTTPException(status_code=400, detail="Unsupported upload type. Use PDF, DOCX, or PPTX.")

        if is_pdf:
            user_content = [
                {
                    "type": "input_text",
                    "text": f"Source file: {file_name}\nExtract structured study information from the attached study description.\nReference data: {json.dumps(payload.get('referenceData', {}))}",
                },
                {"type": "input_file", "filename": file_name, "file_data": file_data},
            ]
        else:
            _, buffer = decode_data_url(file_data)
            extracted = extract_office_text(buffer, "docx" if is_docx else "pptx")
            if not extracted:
                raise HTTPException(status_code=400, detail="The uploaded file did not contain extractable text.")
            user_content = (
                f"Source file: {file_name}\nReference data: {json.dumps(payload.get('referenceData', {}))}\n\n"
                f"Extracted document text:\n{extracted[:120000]}"
            )
        return await call_openai_structured(action, payload, user_content)

    return await call_openai_structured(action, payload)


def clean_text(value: Any, fallback: str = "") -> str:
    text = str(value or fallback)
    text = re.sub(r"\s+\n", "\n", text)
    text = re.sub(r"\n\s+", "\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()


def clean_color(value: Any, fallback: str) -> str:
    color = str(value or "").replace("#", "").strip()
    return color.upper() if re.match(r"^[0-9a-fA-F]{6}$", color) else fallback


def safe_file_name(value: Any) -> str:
    name = re.sub(r"[^\w.\-]+", "-", str(value or "study-schema.pptx"))
    name = re.sub(r"-+", "-", name).strip("-")
    return name if name.endswith(".pptx") else f"{name or 'study-schema'}.pptx"


@app.post("/api/study-schema-pptx")
async def study_schema_pptx(payload: dict[str, Any]) -> Response:
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
    except Exception as exc:
        raise HTTPException(status_code=500, detail="python-pptx is not installed in the backend environment.") from exc

    colors = payload.get("colors", {}) if isinstance(payload.get("colors"), dict) else {}
    accent = clean_color(colors.get("accent"), "D71920")
    edge = clean_color(colors.get("edge"), "8A1538")
    text_color = clean_color(colors.get("text"), "242424")
    background = clean_color(colors.get("background"), "FFF7F7")
    soft_fill = clean_color(colors.get("softFill"), "FDECEC")
    card_fill = clean_color(colors.get("cardFill"), "FFFFFF")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(background)

    def add_text(value: str, x: float, y: float, w: float, h: float, size: float = 10, bold: bool = False, color: str = text_color):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.clear()
        frame.margin_left = Inches(0.03)
        frame.margin_right = Inches(0.03)
        frame.margin_top = Inches(0.02)
        frame.margin_bottom = Inches(0.02)
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.LEFT
        run = paragraph.add_run()
        run.text = clean_text(value)
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
        return box

    def add_card(x: float, y: float, w: float, h: float, fill: str, border: str):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        shape.line.color.rgb = RGBColor.from_string(border)
        shape.line.width = Pt(1.2)
        return shape

    def add_arrow(x1: float, y1: float, x2: float, y2: float):
        connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        connector.line.color.rgb = RGBColor.from_string(edge)
        connector.line.width = Pt(2)
        return connector

    comparator = clean_text(payload.get("comparator"))
    add_text("ONE-SLIDE STUDY SCHEMA", 0.45, 0.18, 3.6, 0.22, 9, True, accent)
    add_card(0.45, 0.42, 12.35, 0.28, soft_fill, soft_fill)
    add_text(clean_text(payload.get("banner"), clean_text(payload.get("designSummary"), "Study design summary")), 0.62, 0.46, 11.95, 0.18, 8.5, True, accent)

    add_card(0.45, 1.45, 2.35, 1.35, accent, accent)
    add_text("STEP 1 · STUDY SETUP", 0.64, 1.64, 1.95, 0.18, 8, True, "FFFFFF")
    add_text(clean_text(payload.get("disease"), "Target condition"), 0.64, 1.94, 1.95, 0.3, 16, True, "FFFFFF")
    add_text(clean_text(payload.get("studySetupDetails")), 0.64, 2.28, 1.95, 0.42, 8.2, False, "FFFFFF")
    add_arrow(2.85, 2.13, 3.25, 2.13)

    add_card(3.32, 1.56, 1.65, 1.2, soft_fill, "F2B8BC")
    add_text("STEP 2", 3.48, 1.8, 0.7, 0.16, 8, True, accent)
    add_text(clean_text(payload.get("designSummary"), "Study design"), 3.48, 2.01, 1.28, 0.5, 10.4, True)
    add_text(clean_text(payload.get("step2Details")), 3.48, 2.52, 1.28, 0.16, 7.4)
    add_arrow(5.0, 2.13, 5.42, 2.13)

    add_card(5.5, 1.6, 2.35, 1.15, accent, accent)
    add_card(5.5, 1.6, 2.35, 0.34, "E5484D", "E5484D")
    add_text("EXPERIMENTAL ARM", 5.66, 1.74, 1.8, 0.16, 8, True, "FFFFFF")
    add_text(clean_text(payload.get("intervention"), "Study intervention"), 5.66, 2.03, 1.9, 0.25, 15, True, "FFFFFF")
    add_text(clean_text(payload.get("interventionSubtitle")), 5.66, 2.38, 1.9, 0.2, 8.5, False, "FFFFFF")

    if comparator:
        add_card(5.5, 2.92, 2.35, 1.08, "F4F4F5", "D4D4D8")
        add_card(5.5, 2.92, 2.35, 0.34, "E5E7EB", "E5E7EB")
        add_text("COMPARATOR ARM", 5.66, 3.06, 1.8, 0.16, 8, True, "52525B")
        add_text(comparator, 5.66, 3.35, 1.9, 0.25, 14, True)
        add_text("Reference treatment strategy", 5.66, 3.68, 1.9, 0.2, 8)
        add_arrow(7.85, 2.18, 8.35, 2.82)
        add_arrow(7.85, 3.46, 8.35, 2.82)
    else:
        add_arrow(7.9, 2.15, 8.45, 2.15)

    add_card(8.55, 1.6, 2.55, 1.35, card_fill, "D6B0BB")
    add_text("STEP 3 · READOUT", 8.75, 1.78, 1.65, 0.18, 8, True, edge)
    add_text(clean_text(payload.get("primaryEndpoint"), "Primary endpoint"), 8.75, 2.1, 2.05, 0.46, 13, True)
    add_text(clean_text(payload.get("evidenceIntent"), "Decision use"), 8.75, 2.62, 2.05, 0.18, 8.5)

    panels = payload.get("panels") if isinstance(payload.get("panels"), list) else []
    for index, panel in enumerate(panels[:6]):
        if not isinstance(panel, dict):
            continue
        row, col = divmod(index, 3)
        x = 0.45 + col * 3.7
        y = 4.35 + row * 1.15
        add_card(x, y, 3.45, 0.88, card_fill, clean_color(panel.get("border"), "F2B8BC"))
        add_text(clean_text(panel.get("title"), "Panel").upper(), x + 0.16, y + 0.12, 3.1, 0.2, 8.5, True, accent)
        add_text(clean_text(panel.get("body"), "To be confirmed"), x + 0.16, y + 0.42, 3.1, 0.38, 8.5, True)

    output = BytesIO()
    prs.save(output)
    return Response(
        output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": f'attachment; filename="{safe_file_name(payload.get("fileName"))}"'},
    )
